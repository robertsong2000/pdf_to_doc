#!/usr/bin/env python3
"""
Service-level fallback detection for pdf2docx conversions.

The fallback is intentionally outside pdf2docx's default behavior: convert once
with the selected mode, inspect the generated DOCX, and only retry with the
fork-only page-frame option when the output looks like a whole-page lattice
table.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from pdf2docx import Converter

from pdf_conversion_modes import CONVERSION_MODE_DEFAULT


PAGE_FRAME_FALLBACK_OPTION = "ignore_page_frame_tables"


def converter_supports_page_frame_fallback(pdf_path: str | Path) -> bool:
    """Return whether the installed pdf2docx build supports the fork option."""
    converter = None
    try:
        converter = Converter(str(pdf_path))
        return PAGE_FRAME_FALLBACK_OPTION in converter.default_settings
    except Exception:
        return False
    finally:
        if converter:
            converter.close()


def page_frame_fallback_options(conversion_options: dict[str, Any]) -> dict[str, Any]:
    options = dict(conversion_options)
    options[PAGE_FRAME_FALLBACK_OPTION] = True
    return options


def should_retry_page_frame_fallback(
    docx_path: str | Path,
    conversion_mode: str | None,
) -> tuple[bool, str]:
    """Detect the pdf2docx whole-page table failure shape.

    This is deliberately conservative: only the default layout mode is retried,
    and the generated DOCX must be dominated by one large table. Real documents
    can contain wide tables, so this should remain a retry heuristic rather than
    a change to pdf2docx defaults.
    """
    if conversion_mode and conversion_mode != CONVERSION_MODE_DEFAULT:
        return False, "fallback only applies to layout mode"

    path = Path(docx_path)
    if not path.exists():
        return False, "output DOCX does not exist"

    document = Document(str(path))
    if not document.tables:
        return False, "output has no tables"

    paragraph_text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )
    paragraph_count = len([p for p in document.paragraphs if p.text.strip()])

    largest_table = max(
        document.tables,
        key=lambda table: len(table.rows) * (len(table.columns) if table.rows else 0),
    )
    rows = len(largest_table.rows)
    cols = len(largest_table.columns) if largest_table.rows else 0
    table_text = "\n".join(
        cell.text.strip()
        for row in largest_table.rows
        for cell in row.cells
        if cell.text.strip()
    )

    if rows < 6 or cols < 6:
        return False, f"largest table is only {rows}x{cols}"

    table_text_len = len(table_text)
    paragraph_text_len = len(paragraph_text)
    table_dominates_document = (
        table_text_len >= 800
        and (paragraph_count <= 5 or table_text_len >= paragraph_text_len * 3)
    )
    if not table_dominates_document:
        return False, (
            f"largest table {rows}x{cols} does not dominate document "
            f"(table chars={table_text_len}, paragraph chars={paragraph_text_len})"
        )

    return True, (
        f"largest table {rows}x{cols} dominates document "
        f"(table chars={table_text_len}, paragraph chars={paragraph_text_len})"
    )
