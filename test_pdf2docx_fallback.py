#!/usr/bin/env python3

import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION

from pdf2docx_fallback import NS, detect_page_frame_table_pages, replace_docx_pages


def _add_large_table(document, prefix: str) -> None:
    table = document.add_table(rows=8, cols=8)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{prefix} cell {row_index}-{col_index} " + ("content " * 4)
    table.cell(0, 0).text = "Document Name"
    table.cell(1, 0).text = "Verification Method"
    table.cell(2, 0).text = "Legacy ID"


class PageFrameFallbackTests(unittest.TestCase):
    def test_detects_only_page_segment_with_dominating_table(self):
        document = Document()
        document.add_paragraph("Page one normal paragraph")
        document.add_section(WD_SECTION.NEW_PAGE)
        _add_large_table(document, "bad page")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("Page three normal paragraph")

        with tempfile.TemporaryDirectory() as tmp_dir:
            docx_path = Path(tmp_dir) / "page-frame.docx"
            document.save(docx_path)

            pages, reason = detect_page_frame_table_pages(docx_path, "layout")

            self.assertEqual(pages, [1])
            self.assertIn("page 2", reason)

    def test_replace_docx_pages_keeps_unmatched_pages(self):
        target = Document()
        target.add_paragraph("page one stays")
        target.add_section(WD_SECTION.NEW_PAGE)
        _add_large_table(target, "bad page")
        target.add_section(WD_SECTION.NEW_PAGE)
        target.add_paragraph("page three stays")

        replacement = Document()
        replacement.add_paragraph("page two replacement")

        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "target.docx"
            replacement_path = Path(tmp_dir) / "replacement.docx"
            target.save(target_path)
            replacement.save(replacement_path)

            replace_docx_pages(target_path, {1: replacement_path})

            cleaned = Document(str(target_path))
            text = "\n".join(paragraph.text for paragraph in cleaned.paragraphs)
            self.assertIn("page one stays", text)
            self.assertIn("page two replacement", text)
            self.assertIn("page three stays", text)
            self.assertNotIn("bad page cell", text)

    def test_replace_docx_pages_preserves_target_page_boundary(self):
        target = Document()
        target.add_paragraph("page one stays")
        target.add_section(WD_SECTION.NEW_PAGE)
        _add_large_table(target, "bad page")
        target.add_section(WD_SECTION.NEW_PAGE)
        target.add_paragraph("page three stays")

        replacement = Document()
        replacement.add_paragraph("page two replacement")

        with tempfile.TemporaryDirectory() as tmp_dir:
            target_path = Path(tmp_dir) / "target.docx"
            replacement_path = Path(tmp_dir) / "replacement.docx"
            target.save(target_path)
            replacement.save(replacement_path)

            with zipfile.ZipFile(target_path, "r") as docx_zip:
                target_root = etree.fromstring(docx_zip.read("word/document.xml"))
            target_body = target_root.find("w:body", NS)
            original_boundary = etree.tostring(list(target_body)[3])

            replace_docx_pages(target_path, {1: replacement_path})

            with zipfile.ZipFile(target_path, "r") as docx_zip:
                cleaned_root = etree.fromstring(docx_zip.read("word/document.xml"))
            cleaned_body = cleaned_root.find("w:body", NS)
            cleaned_children = list(cleaned_body)

            self.assertEqual(etree.tostring(cleaned_children[3]), original_boundary)
            self.assertNotEqual(cleaned_children[3].tag, f"{{{NS['w']}}}sectPr")


if __name__ == "__main__":
    unittest.main()
