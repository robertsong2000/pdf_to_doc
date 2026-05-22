#!/usr/bin/env python3

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_header_cleaner import remove_repeated_pdf_headers


def _set_vertical_merge(cell, value: str | None) -> None:
    properties = cell._tc.get_or_add_tcPr()
    vertical_merge = OxmlElement("w:vMerge")
    if value is not None:
        vertical_merge.set(qn("w:val"), value)
    properties.append(vertical_merge)


class HeaderCleanupTests(unittest.TestCase):
    def test_header_cleanup_removes_orphan_vertical_merges(self):
        document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Geely Automotive Research Institute"
        table.cell(0, 1).text = "Document Type Document Release Status"
        table.cell(1, 0).text = "Document No Revision"
        table.cell(1, 1).text = "Volume No Page No"
        table.cell(2, 0).text = "Document Name"
        table.cell(2, 1).text = "Body content after the header"

        _set_vertical_merge(table.cell(0, 0), "restart")
        _set_vertical_merge(table.cell(1, 0), None)
        _set_vertical_merge(table.cell(2, 0), None)

        with tempfile.TemporaryDirectory() as tmp_dir:
            docx_path = Path(tmp_dir) / "header-with-merge.docx"
            document.save(docx_path)

            removed = remove_repeated_pdf_headers(docx_path)

            self.assertEqual(removed, 2)
            cleaned = Document(str(docx_path))
            self.assertEqual(len(cleaned.tables[0].rows), 1)
            self.assertEqual(cleaned.tables[0].rows[0].cells[0].text, "Document Name")
            self.assertEqual(
                cleaned.tables[0].rows[0].cells[1].text,
                "Body content after the header",
            )

    def test_header_cleanup_removes_leading_header_paragraphs(self):
        document = Document()
        for text in (
            "Document Type Document Release Status",
            "Geely Automotive Research Institute",
            "Document No",
            "x04000000 Revision",
            "001 Volume No Page No",
            "1886(2573) NOTE-SWRS RELEASED",
            "(Ningbo)Co.Ltd 184C0769",
            "Document Name",
            "SWRS-ZCUDM_V01_23R3U3_GEEA3.0",
            "146582 v1 General VMM requirement [CS Released]",
        ):
            document.add_paragraph(text)

        with tempfile.TemporaryDirectory() as tmp_dir:
            docx_path = Path(tmp_dir) / "header-paragraphs.docx"
            document.save(docx_path)

            removed = remove_repeated_pdf_headers(docx_path)

            self.assertEqual(removed, 7)
            cleaned = Document(str(docx_path))
            remaining = [p.text for p in cleaned.paragraphs if p.text.strip()]
            self.assertEqual(remaining[0], "Document Name")
            self.assertEqual(remaining[1], "SWRS-ZCUDM_V01_23R3U3_GEEA3.0")
            self.assertEqual(
                remaining[2],
                "146582 v1 General VMM requirement [CS Released]",
            )


if __name__ == "__main__":
    unittest.main()
